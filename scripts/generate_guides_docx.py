from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"


def set_font(run, size=None, bold=None, color=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def border_table(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "D9D9D9")
        borders.append(e)
    tbl_pr.append(borders)


def add_text(p, text, size=10.5, bold=False):
    for i, part in enumerate(re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)):
        if not part:
            continue
        if part.startswith("`"):
            run = p.add_run(part[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(9.5)
        elif part.startswith("**"):
            run = p.add_run(part[2:-2]); set_font(run, size=size, bold=True)
        else:
            run = p.add_run(part); set_font(run, size=size, bold=bold)


def new_document(title, purpose):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.1); section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.25); section.right_margin = Cm(2.25)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"; styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size in (("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)):
        st = styles[style_name]; st.font.name = "Aptos"; st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        st.font.size = Pt(size); st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.bold = True
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title); set_font(run, size=22, bold=True, color=(0, 0, 0))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, purpose, size=11)
    p.paragraph_format.space_after = Pt(18)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "Voca Basic", size=9)
    return doc


def add_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"; border_table(table)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for col, value in enumerate(row):
            p = cells[col].paragraphs[0]
            add_text(p, value, size=9.5, bold=(index == 0))
            if index == 0:
                shade(cells[col], "D9EAF7")
            elif index % 2 == 0:
                shade(cells[col], "F7FBFE")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(markdown_path, output_path, title, purpose):
    doc = new_document(title, purpose)
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    in_code = False; code_lines = []; table_rows = []

    def flush_code():
        nonlocal code_lines
        if code_lines:
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run("\n".join(code_lines)); run.font.name = "Consolas"; run.font.size = Pt(8.5)
            code_lines = []

    def flush_table():
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows); table_rows = []

    for line in lines:
        if line.startswith("```"):
            in_code = not in_code
            if not in_code: flush_code()
            continue
        if in_code:
            code_lines.append(line); continue
        if line.startswith("|") and not line.startswith("| ---"):
            table_rows.append([cell.strip() for cell in line.strip("|").split("|")])
            continue
        flush_table()
        if not line.strip():
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(line[3:], style="Heading 1"); p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
        elif line.startswith("### "):
            p = doc.add_paragraph(line[4:], style="Heading 2"); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number"); add_text(p, re.sub(r"^\d+\. ", "", line))
            p.paragraph_format.space_after = Pt(3)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); add_text(p, line[2:])
            p.paragraph_format.space_after = Pt(3)
        else:
            p = doc.add_paragraph(); add_text(p, line)
            p.paragraph_format.space_after = Pt(6)
    flush_code(); flush_table()
    doc.core_properties.title = title
    doc.core_properties.subject = purpose
    doc.core_properties.author = "Voca Basic"
    doc.save(output_path)


convert(OUT / "HUONG_DAN_CAI_DAT.md", OUT / "Huong Dan Cai Dat Voca Basic.docx", "Hướng dẫn cài đặt Voca Basic", "Cài đặt Core, gói AI và chuẩn bị máy để sử dụng Voca Basic")
convert(OUT / "HUONG_DAN_SU_DUNG.md", OUT / "Huong Dan Su Dung Voca Basic.docx", "Hướng dẫn sử dụng Voca Basic", "Tạo giọng nói, nhân bản giọng và quản lý dữ liệu trong Voca Basic")

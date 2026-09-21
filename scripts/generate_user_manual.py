"""Script to generate comprehensive User Manual files for Voca Basic in TXT, DOCX, and HTML formats."""
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def build_docx(out_path: Path):
    doc = docx.Document()

    # Page setup - normal margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("VOCA BASIC — HƯỚNG DẪN CÀI ĐẶT & SỬ DỤNG")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138)  # Navy Blue

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Phần Mềm Thuyết Minh & Nhân Bản Giọng Nói AI Offline 100%\nPhiên bản Thương mại 2026")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # Section 1: Giới thiệu
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. GIỚI THIỆU PHẦN MỀM")
    r.font.color.rgb = RGBColor(37, 99, 235)
    
    doc.add_paragraph(
        "Voca Basic là giải pháp phần mềm thuyết minh, chuyển đổi văn bản thành giọng nói (Text-To-Speech) "
        "và nhân bản giọng nói (Voice Cloning) bằng trí tuệ nhân tạo thế hệ mới. "
        "Phần mềm hoạt động hoàn toàn Offline trên máy tính cá nhân của bạn, không cần kết nối Internet, "
        "bảo mật tuyệt đối dữ liệu và không giới hạn số lượng ký tự chuyển đổi."
    )

    doc.add_paragraph("Các tính năng nổi bật:")
    features = [
        "Chuyển văn bản thành giọng đọc tự nhiên (tiếng Việt chuẩn các vùng miền Bắc - Nam, tiếng Anh, tiếng Trung...).",
        "Nhân bản giọng nói AI (Voice Cloning): Chỉ cần mẫu ghi âm 5 - 15 giây, AI sẽ đọc theo đúng ngữ điệu và màu giọng của bạn.",
        "Đọc văn bản dài (Long-form): Tự động ngắt nhịp và xử lý các bài thuyết minh, sách nói, kịch bản dài mượt mà.",
        "Tùy chỉnh linh hoạt: Điều chỉnh tốc độ (Speed), độ cao (Pitch), xuất file âm thanh chất lượng cao MP3 / WAV.",
        "Bản quyền bảo mật: Xác thực mã máy Offline, không lo mất bản quyền khi mất mạng."
    ]
    for feat in features:
        p = doc.add_paragraph(feat, style='List Bullet')
        p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph()

    # Section 2: Hướng dẫn cài đặt
    h2 = doc.add_heading(level=1)
    r2 = h2.add_run("2. HƯỚNG DẪN CÀI ĐẶT & GHÉP NỐI DỮ LIỆU (LẦN ĐẦU TIÊN)")
    r2.font.color.rgb = RGBColor(37, 99, 235)

    doc.add_paragraph(
        "Để tối ưu tốc độ tải về, bộ cài đặt được chia thành 2 phần: Bản Core nhẹ (~130 MB) và Bộ gói dữ liệu AI (~8.8 GB). "
        "Bạn chỉ cần thực hiện ghép nối theo các bước sau một lần duy nhất:"
    )

    steps = [
        ("Bước 1: Chuẩn bị thư mục chứa phần mềm", 
         "Giải nén file Voca-Basic-Core.zip vào một ổ đĩa có dung lượng trống tối thiểu 20 GB (khuyên dùng ổ D: hoặc E:, ví dụ: D:\\Voca-Basic). Tránh để ở Desktop hoặc ổ C: nếu ổ C bị đầy."),
        ("Bước 2: Tạo thư mục ai-packages",
         "Mở thư mục Voca Basic Core vừa giải nén, tạo một thư mục mới đặt tên chính xác là: ai-packages (nằm cùng cấp với file Voca Basic.exe)."),
        ("Bước 3: Giải nén 3 gói AI vào thư mục ai-packages",
         "Dùng phần mềm 7-Zip hoặc WinRAR, giải nén 3 file .7z được cấp vào đúng thư mục ai-packages vừa tạo:\n"
         "  • runtime-main.7z  -->  ai-packages\\runtime-main\\\n"
         "  • runtime-chatterbox.7z  -->  ai-packages\\runtime-chatterbox\\\n"
         "  • models.7z  -->  ai-packages\\models\\"),
        ("Bước 4: Tạo lối tắt ngoài màn hình chính",
         "Nhấp đúp vào file Tao-Phim-Tat-Desktop.bat trong thư mục để tự động tạo biểu tượng 'Voca Basic' ngoài Desktop."),
        ("Bước 5: Khởi động phần mềm",
         "Nhấp vào biểu tượng ngoài Desktop hoặc chạy file Voca Basic.exe để mở ứng dụng.")
    ]
    for title, desc in steps:
        p = doc.add_paragraph()
        r_t = p.add_run(f"• {title}: ")
        r_t.bold = True
        p.add_run(desc)
        p.paragraph_format.space_after = Pt(4)

    doc.add_paragraph()

    # Section 3: Kích hoạt bản quyền
    h3 = doc.add_heading(level=1)
    r3 = h3.add_run("3. HƯỚNG DẪN KÍCH HOẠT BẢN QUYỀN")
    r3.font.color.rgb = RGBColor(37, 99, 235)

    lic_steps = [
        ("1. Lấy Mã máy (Machine ID)", "Khi mở ứng dụng lần đầu, góc trên bên phải màn hình sẽ có huy hiệu màu đỏ '⚠️ Chưa kích hoạt bản quyền'. Bạn nhấp vào đó, bấm nút '📋 Sao chép mã máy' (Mã có dạng: VB-XXXX-XXXX-XXXX-XXXX)."),
        ("2. Gửi mã máy cho người bán", "Dán mã máy vừa copy vào tin nhắn Zalo/Telegram và gửi cho người bán để nhận mã bản quyền (License Key)."),
        ("3. Kích hoạt vào phần mềm", "Copy chuỗi License Key nhận được, dán vào ô 'Nhập mã bản quyền' trong ứng dụng và bấm nút '🚀 Kích hoạt ngay'."),
        ("4. Hoàn tất", "Huy hiệu sẽ chuyển sang màu xanh lá '🛡️ Bản quyền: Vĩnh viễn' (hoặc hiển thị thời gian còn lại nếu dùng gói có thời hạn). Tất cả tính năng tạo giọng và nhân bản đều được mở khóa ngay.")
    ]
    for title, desc in lic_steps:
        p = doc.add_paragraph()
        r_t = p.add_run(f"{title}: ")
        r_t.bold = True
        p.add_run(desc)

    doc.add_paragraph()

    # Section 4: Hướng dẫn sử dụng tính năng
    h4 = doc.add_heading(level=1)
    r4 = h4.add_run("4. HƯỚNG DẪN SỬ DỤNG CÁC TÍNH NĂNG CHÍNH")
    r4.font.color.rgb = RGBColor(37, 99, 235)

    p = doc.add_paragraph()
    r = p.add_run("4.1. Chuyển Văn Bản Thành Giọng Nói (Text-to-Speech)")
    r.bold = True
    r.font.size = Pt(12)
    doc.add_paragraph(
        "1. Chọn thẻ 'Tạo giọng nói' trên thanh menu bên trái.\n"
        "2. Nhập hoặc dán đoạn văn bản cần thuyết minh vào khung soạn thảo.\n"
        "3. Chọn giọng đọc trong danh sách (giọng nam/nữ, Bắc/Nam, hoặc giọng nhân bản của riêng bạn).\n"
        "4. Kéo thanh trượt để điều chỉnh Tốc độ đọc (Speed) hoặc Âm lượng (Volume) theo ý muốn.\n"
        "5. Nhấp nút 'Tạo âm thanh': AI sẽ xử lý và cho bạn nghe thử trực tiếp.\n"
        "6. Bấm nút 'Tải về' để lưu file âm thanh dưới dạng MP3 hoặc WAV về máy tính."
    )

    p = doc.add_paragraph()
    r = p.add_run("4.2. Đọc Văn Bản Dài (Long-Form TTS)")
    r.bold = True
    r.font.size = Pt(12)
    doc.add_paragraph(
        "Đối với các bài báo, truyện ngắn, hoặc kịch bản video dài:\n"
        "• Dán toàn bộ nội dung văn bản vào khung 'Văn bản dài'.\n"
        "• Hệ thống tự động phân tách từng câu/đoạn văn bản một cách thông minh, đảm bảo âm thanh đọc liền mạch, không bị mất từ hay giật lag.\n"
        "• Cho phép xem trước danh sách từng câu và xuất toàn bộ ra một file âm thanh hoàn chỉnh."
    )

    p = doc.add_paragraph()
    r = p.add_run("4.3. Nhân Bản Giọng Nói AI (Voice Cloning)")
    r.bold = True
    r.font.size = Pt(12)
    doc.add_paragraph(
        "Tính năng độc quyền giúp AI bắt chước chính xác chất giọng của bạn hoặc người khác:\n"
        "1. Chuẩn bị file ghi âm: Ghi âm giọng nói mẫu từ 5 đến 15 giây (file .mp3 hoặc .wav). Yêu cầu: đọc tự nhiên, rõ tiếng, không có nhạc nền hay tiếng ồn xung quanh.\n"
        "2. Vào mục 'Nhân bản giọng nói' (Voice Cloning) trên menu.\n"
        "3. Nhấp 'Tải lên mẫu âm thanh' và chọn file ghi âm của bạn.\n"
        "4. Đặt tên gợi nhớ cho hồ sơ giọng (ví dụ: 'Giọng đọc của tôi', 'MC Nam').\n"
        "5. Bấm 'Tạo hồ sơ giọng'. AI sẽ học biểu cảm và màu giọng chỉ trong vài giây.\n"
        "6. Giờ đây, trong mục 'Tạo giọng nói', bạn chỉ cần chọn tên hồ sơ này là AI sẽ đọc bất kỳ văn bản nào bằng chính giọng của bạn!"
    )

    doc.add_paragraph()

    # Section 5: Khắc phục sự cố
    h5 = doc.add_heading(level=1)
    r5 = h5.add_run("5. KHẮC PHỤC SỰ CỐ & CÂU HỎI THƯỜNG GẶP")
    r5.font.color.rgb = RGBColor(37, 99, 235)

    faqs = [
        ("Lỗi: 'Missing packaged runtime-main Python...'",
         "Nguyên nhân do bạn chưa giải nén đủ 3 gói AI hoặc thư mục 'ai-packages' đặt sai vị trí. Hãy đảm bảo thư mục 'ai-packages' nằm ngay cạnh file 'Voca Basic.exe' và bên trong có đủ 3 thư mục: runtime-main, runtime-chatterbox, models."),
        ("Làm sao để đưa ứng dụng về bản trắng sạch ban đầu?",
         "Trong thư mục ứng dụng có file 'Reset-Data-Ban-Trang.bat'. Bạn chỉ cần tắt app và click đúp vào file này, toàn bộ lịch sử và bản quyền cũ trên máy sẽ được dọn sạch hoàn toàn."),
        ("Bản quyền có cần kết nối Internet để kiểm tra định kỳ không?",
         "Hoàn toàn không. Voca Basic sử dụng công nghệ ký số mật mã Ed25519 cục bộ, xác thực 100% Offline ngay trên máy tính của bạn."),
        ("Yêu cầu cấu hình máy tính khuyến nghị?",
         "Hệ điều hành: Windows 10 hoặc Windows 11 (64-bit).\n"
         "RAM: Tối thiểu 8 GB (Khuyến nghị 16 GB để AI chạy mượt mà nhất).\n"
         "Ổ cứng trống: Tối thiểu 20 GB (ưu tiên ổ SSD để tải mô hình nhanh chóng).")
    ]
    for q, a in faqs:
        p = doc.add_paragraph()
        r_q = p.add_run(f"❓ {q}\n")
        r_q.bold = True
        r_q.font.color.rgb = RGBColor(15, 23, 42)
        p.add_run(f"👉 Trả lời: {a}")
        p.paragraph_format.space_after = Pt(6)

    # Footer note
    doc.add_paragraph()
    p_f = doc.add_paragraph()
    p_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_f = p_f.add_run("— CHÚC BẠN CÓ NHỮNG TRẢI NGHIỆM TUYỆT VỜI CÙNG VOCA BASIC! —")
    r_f.font.name = "Segoe UI"
    r_f.font.bold = True
    r_f.font.color.rgb = RGBColor(71, 85, 105)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Saved DOCX: {out_path}")


def build_txt(out_path: Path):
    content = """================================================================================
               VOCA BASIC — HƯỚNG DẪN CÀI ĐẶT & SỬ DỤNG
          Phần Mềm Thuyết Minh & Nhân Bản Giọng Nói AI Offline 100%
================================================================================

1. GIỚI THIỆU PHẦN MỀM
--------------------------------------------------------------------------------
Voca Basic là giải pháp chuyển đổi văn bản thành giọng nói (Text-To-Speech) và
nhân bản giọng nói (Voice Cloning) thế hệ mới bằng trí tuệ nhân tạo (AI).
- Hoạt động 100% Offline trên máy tính, không cần kết nối mạng.
- Bảo mật dữ liệu tuyệt đối, không giới hạn ký tự.
- Giọng đọc tự nhiên, chuẩn truyền cảm tiếng Việt, tiếng Anh, tiếng Trung...
- Nhân bản chất giọng bất kỳ chỉ với 5-15 giây file ghi âm mẫu.


2. HƯỚNG DẪN CÀI ĐẶT & GHÉP NỐI DỮ LIỆU (LẦN ĐẦU TIÊN)
--------------------------------------------------------------------------------
Bộ cài đặt gồm 2 phần:
  - Bản Core nhẹ: Voca-Basic-Core.zip (~130 MB)
  - Bộ 3 gói AI: runtime-main.7z, runtime-chatterbox.7z, models.7z (~8.8 GB)

Các bước cài đặt:
• BƯỚC 1: Giải nén file "Voca-Basic-Core.zip" vào một thư mục trên ổ đĩa
          (Khuyên dùng ổ D: hoặc E:, ví dụ: D:\\Voca-Basic).
• BƯỚC 2: Vào trong thư mục vừa giải nén, tạo một thư mục mới tên là:
          ai-packages (nằm ngay cạnh file Voca Basic.exe).
• BƯỚC 3: Dùng phần mềm 7-Zip hoặc WinRAR giải nén 3 file .7z vào thư mục
          ai-packages vừa tạo:
          + runtime-main.7z       -->  ai-packages\\runtime-main\\
          + runtime-chatterbox.7z -->  ai-packages\\runtime-chatterbox\\
          + models.7z             -->  ai-packages\\models\\
• BƯỚC 4: Nhấp đúp vào file "Tao-Phim-Tat-Desktop.bat" để tạo biểu tượng mở
          nhanh ngoài màn hình Desktop.
• BƯỚC 5: Mở "Voca Basic.exe" để bắt đầu sử dụng.


3. HƯỚNG DẪN KÍCH HOẠT BẢN QUYỀN
--------------------------------------------------------------------------------
• BƯỚC 1: Mở ứng dụng Voca Basic.
• BƯỚC 2: Nhấp vào huy hiệu màu đỏ "⚠️ Chưa kích hoạt bản quyền" ở góc trên bên
          phải màn hình, bấm nút "📋 Sao chép mã máy" (dạng: VB-XXXX-XXXX-XXXX-XXXX).
• BƯỚC 3: Gửi mã máy này cho người bán qua Zalo/Telegram.
• BƯỚC 4: Khi nhận được License Key từ người bán, dán vào ô "Nhập mã bản quyền"
          trong ứng dụng và bấm nút "🚀 Kích hoạt ngay".
• BƯỚC 5: Ứng dụng thông báo kích hoạt thành công, huy hiệu chuyển sang màu xanh
          lá và mở khóa toàn bộ tính năng sử dụng trọn đời hoặc theo hạn đăng ký!


4. HƯỚNG DẪN SỬ DỤNG CÁC TÍNH NĂNG CHÍNH
--------------------------------------------------------------------------------
4.1. Chuyển văn bản thành giọng nói (Text-to-Speech):
  - Chọn mục "Tạo giọng nói" trên thanh menu bên trái.
  - Nhập hoặc dán đoạn văn bản cần đọc vào ô soạn thảo.
  - Chọn giọng đọc mong muốn (Nam/Nữ, Bắc/Nam, hoặc giọng nhân bản).
  - Tùy chỉnh tốc độ (Speed) và âm lượng nếu cần.
  - Nhấp "Tạo âm thanh" để nghe thử, bấm "Tải về" để lưu file .MP3 / .WAV.

4.2. Thuyết minh văn bản dài (Long-Form TTS):
  - Dán bài viết, truyện ngắn, hoặc kịch bản dài nhiều trang vào mục "Văn bản dài".
  - Hệ thống tự động tách câu thông minh, đọc mượt mà không bị ngắt quãng.
  - Xuất ra 1 file âm thanh hoàn chỉnh cho toàn bộ bài thuyết minh.

4.3. Nhân bản giọng nói AI (Voice Cloning):
  - Chuẩn bị file ghi âm giọng nói mẫu (5 - 15 giây), định dạng .mp3 hoặc .wav.
    (Lưu ý: Nói tự nhiên, to rõ, không có nhạc nền hay tiếng ồn).
  - Vào mục "Nhân bản giọng nói" (Voice Cloning) trên menu.
  - Tải file ghi âm lên, đặt tên cho giọng đọc (ví dụ: Giọng Anh Nam).
  - Bấm "Tạo hồ sơ giọng". AI sẽ học chất giọng chỉ trong vài giây.
  - Vào lại mục "Tạo giọng nói", chọn tên hồ sơ vừa tạo để AI nói bằng chính
    giọng của bạn!


5. KHẮC PHỤC SỰ CỐ & CÂU HỎI THƯỜNG GẶP
--------------------------------------------------------------------------------
❓ Lỗi "Missing packaged runtime-main Python (kiểm tra thư mục runtime-main/python.exe)":
   👉 Trả lời: Do bạn chưa giải nén đủ 3 gói AI hoặc đặt sai vị trí thư mục.
      Đảm bảo thư mục "ai-packages" nằm ngay cùng cấp với "Voca Basic.exe"
      và bên trong có đủ 3 thư mục con: runtime-main, runtime-chatterbox, models.

❓ Làm sao để reset phần mềm về bản trắng sạch sẽ?
   👉 Trả lời: Click đúp vào file "Reset-Data-Ban-Trang.bat" trong thư mục ứng
      dụng để xóa sạch toàn bộ lịch sử và đưa app về trạng thái trắng ban đầu.

❓ Yêu cầu cấu hình máy tính:
   👉 Trả lời: Windows 10/11 (64-bit), RAM tối thiểu 8GB (khuyên dùng 16GB),
      ổ cứng trống tối thiểu 20GB.

================================================================================
          CẢM ƠN BẠN ĐÃ TIN TƯỞNG VÀ SỬ DỤNG VOCA BASIC!
================================================================================
"""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(content, encoding="utf-8")
    print(f"Saved TXT: {out_path}")


def build_html(out_path: Path):
    html = """<!DOCTYPE html>
<html lang="vi">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Voca Basic — Hướng Dẫn Cài Đặt & Sử Dụng</title>
  <style>
    :root {
      --primary: #2563eb;
      --primary-dark: #1d4ed8;
      --bg: #f8fafc;
      --card-bg: #ffffff;
      --text: #1e293b;
      --text-muted: #64748b;
      --border: #e2e8f0;
      --success: #16a34a;
      --warning: #f59e0b;
    }
    * { box-sizing: border-box; margin: 0; padding: 0; }
    body {
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
      background-color: var(--bg);
      color: var(--text);
      line-height: 1.6;
      padding: 30px 15px;
    }
    .container {
      max-width: 860px;
      margin: 0 auto;
      background: var(--card-bg);
      border-radius: 12px;
      box-shadow: 0 4px 6px -1px rgba(0,0,0,0.05), 0 2px 4px -2px rgba(0,0,0,0.05);
      border: 1px solid var(--border);
      overflow: hidden;
    }
    header {
      background: linear-gradient(135deg, #1e3a8a, #2563eb);
      color: #fff;
      padding: 36px 30px;
      text-align: center;
    }
    header h1 { font-size: 26px; font-weight: 800; letter-spacing: 0.5px; margin-bottom: 8px; }
    header p { font-size: 15px; opacity: 0.9; }
    .content { padding: 32px 30px; }
    h2 {
      font-size: 19px;
      color: var(--primary-dark);
      border-bottom: 2px solid #dbeafe;
      padding-bottom: 8px;
      margin-top: 32px;
      margin-bottom: 16px;
      display: flex;
      align-items: center;
      gap: 8px;
    }
    h2:first-of-type { margin-top: 0; }
    h3 { font-size: 16px; color: #0f172a; margin: 18px 0 8px; }
    p, li { font-size: 15px; margin-bottom: 10px; color: #334155; }
    ul, ol { padding-left: 22px; margin-bottom: 16px; }
    li { margin-bottom: 6px; }
    .step-card {
      background: #f1f5f9;
      border-left: 4px solid var(--primary);
      padding: 14px 18px;
      border-radius: 0 8px 8px 0;
      margin-bottom: 14px;
    }
    .step-title { font-weight: 700; color: #0f172a; margin-bottom: 4px; }
    .alert-box {
      background: #eff6ff;
      border: 1px solid #bfdbfe;
      border-radius: 8px;
      padding: 16px;
      margin: 18px 0;
    }
    .alert-warning {
      background: #fffbeb;
      border-color: #fde68a;
    }
    .code-tag {
      background: #e2e8f0;
      padding: 2px 6px;
      border-radius: 4px;
      font-family: Consolas, monospace;
      font-size: 13.5px;
      font-weight: bold;
      color: #0f172a;
    }
    .badge {
      display: inline-block;
      padding: 3px 10px;
      border-radius: 9999px;
      font-size: 12px;
      font-weight: 700;
      color: white;
    }
    .badge-success { background: var(--success); }
    .faq-item {
      background: #f8fafc;
      border: 1px solid var(--border);
      border-radius: 8px;
      padding: 14px 18px;
      margin-bottom: 12px;
    }
    .faq-q { font-weight: 700; color: #0f172a; margin-bottom: 6px; }
    footer {
      background: #f1f5f9;
      text-align: center;
      padding: 20px;
      font-size: 14px;
      color: var(--text-muted);
      border-top: 1px solid var(--border);
    }
  </style>
</head>
<body>

<div class="container">
  <header>
    <h1>VOCA BASIC — HƯỚNG DẪN CÀI ĐẶT & SỬ DỤNG</h1>
    <p>Phần Mềm Thuyết Minh & Nhân Bản Giọng Nói AI Offline 100%</p>
  </header>

  <div class="content">
    <h2>1. Giới Thiệu Phần Mềm</h2>
    <p><strong>Voca Basic</strong> là ứng dụng chuyển đổi văn bản thành giọng nói (Text-To-Speech) và nhân bản giọng nói (Voice Cloning) bằng trí tuệ nhân tạo (AI) chạy hoàn toàn <strong>Offline</strong> trên máy tính của bạn.</p>
    <ul>
      <li><strong>Hoạt động 100% Offline:</strong> Không cần Internet, bảo mật tuyệt đối dữ liệu cá nhân.</li>
      <li><strong>Không giới hạn ký tự:</strong> Thoải mái đọc các bài văn, truyện, sách nói dài.</li>
      <li><strong>Chất giọng tự nhiên:</strong> Hỗ trợ đa dạng giọng đọc truyền cảm các miền Bắc - Nam tiếng Việt, tiếng Anh, tiếng Trung...</li>
      <li><strong>Nhân bản giọng nói (Voice Cloning):</strong> Chỉ cần mẫu ghi âm 5 - 15 giây, AI sẽ nói theo đúng chất giọng của bạn!</li>
    </ul>

    <h2>2. Hướng Dẫn Cài Đặt & Ghép Nối Dữ Liệu (Lần đầu tiên)</h2>
    <p>Bộ cài đặt được chia thành <strong>Bản Core nhẹ (~130 MB)</strong> và <strong>Bộ 3 gói AI (~8.8 GB)</strong> để thuận tiện tải về:</p>

    <div class="step-card">
      <div class="step-title">Bước 1: Chuẩn bị thư mục</div>
      <p>Giải nén file <span class="code-tag">Voca-Basic-Core.zip</span> vào một ổ đĩa có dung lượng trống từ 20 GB trở lên (Khuyên dùng ổ D: hoặc E:, ví dụ: <code>D:\\Voca-Basic</code>).</p>
    </div>

    <div class="step-card">
      <div class="step-title">Bước 2: Tạo thư mục ai-packages</div>
      <p>Vào thư mục vừa giải nén, tạo một thư mục mới đặt tên chính xác là: <span class="code-tag">ai-packages</span> (nằm cùng cấp với file <code>Voca Basic.exe</code>).</p>
    </div>

    <div class="step-card">
      <div class="step-title">Bước 3: Giải nén 3 gói AI vào ai-packages</div>
      <p>Dùng 7-Zip hoặc WinRAR giải nén 3 file <span class="code-tag">.7z</span> được cấp vào đúng thư mục <code>ai-packages</code>:</p>
      <ul>
        <li><code>runtime-main.7z</code> &rarr; thư mục <code>ai-packages\\runtime-main\\</code></li>
        <li><code>runtime-chatterbox.7z</code> &rarr; thư mục <code>ai-packages\\runtime-chatterbox\\</code></li>
        <li><code>models.7z</code> &rarr; thư mục <code>ai-packages\\models\\</code></li>
      </ul>
    </div>

    <div class="step-card">
      <div class="step-title">Bước 4: Tạo lối tắt ngoài Desktop</div>
      <p>Click đúp file <span class="code-tag">Tao-Phim-Tat-Desktop.bat</span> để tự động tạo biểu tượng ứng dụng ngoài màn hình chính.</p>
    </div>

    <div class="step-card">
      <div class="step-title">Bước 5: Khởi động</div>
      <p>Click biểu tượng ngoài Desktop hoặc chạy file <span class="code-tag">Voca Basic.exe</span> để bắt đầu.</p>
    </div>

    <h2>3. Hướng Dẫn Kích Hoạt Bản Quyền</h2>
    <div class="alert-box">
      <p><strong>1. Lấy Mã máy (Machine ID):</strong> Mở ứng dụng, nhìn góc trên bên phải có huy hiệu màu đỏ <span class="code-tag">⚠️ Chưa kích hoạt bản quyền</span>. Nhấp vào đó và bấm <strong>"📋 Sao chép mã máy"</strong> (Mã có dạng: <code>VB-XXXX-XXXX-XXXX-XXXX</code>).</p>
      <p><strong>2. Gửi mã máy cho người bán:</strong> Gửi mã này qua Zalo/Telegram để nhận chuỗi <strong>License Key</strong>.</p>
      <p><strong>3. Dán key & kích hoạt:</strong> Copy chuỗi License Key nhận được, dán vào ô nhập trong ứng dụng và bấm <strong>"🚀 Kích hoạt ngay"</strong>.</p>
      <p><strong>4. Hoàn tất:</strong> Huy hiệu sẽ chuyển sang màu xanh lá <span class="badge badge-success">🛡️ Bản quyền: Vĩnh viễn</span> và mở khóa toàn bộ tính năng sử dụng ngay lập tức!</p>
    </div>

    <h2>4. Hướng Dẫn Sử Dụng Các Tính Năng Chính</h2>
    <h3>4.1. Chuyển văn bản thành giọng nói (Text-to-Speech)</h3>
    <ol>
      <li>Chọn mục <strong>"Tạo giọng nói"</strong> trên thanh menu bên trái.</li>
      <li>Nhập hoặc dán đoạn văn bản cần thuyết minh vào khung soạn thảo.</li>
      <li>Chọn giọng đọc mong muốn trong danh sách.</li>
      <li>Điều chỉnh <strong>Tốc độ đọc (Speed)</strong> hoặc âm lượng nếu muốn.</li>
      <li>Bấm <strong>"Tạo âm thanh"</strong> để nghe thử, bấm <strong>"Tải về"</strong> để lưu file MP3 / WAV.</li>
    </ol>

    <h3>4.2. Thuyết minh văn bản dài (Long-Form)</h3>
    <p>Đối với kịch bản video, bài báo, sách nói nhiều trang:</p>
    <ul>
      <li>Dán toàn bộ nội dung vào mục <strong>"Văn bản dài"</strong>.</li>
      <li>Hệ thống tự động tách câu thông minh, đọc mượt mà không bị ngắt chữ hay gián đoạn.</li>
      <li>Xuất toàn bộ bài thuyết minh ra một file âm thanh hoàn chỉnh.</li>
    </ul>

    <h3>4.3. Nhân bản giọng nói AI (Voice Cloning)</h3>
    <ol>
      <li><strong>Chuẩn bị mẫu giọng:</strong> Ghi âm giọng nói mẫu từ 5 đến 15 giây (file .mp3 hoặc .wav), nói rõ tiếng, không có nhạc nền.</li>
      <li>Vào mục <strong>"Nhân bản giọng nói" (Voice Cloning)</strong> trên menu.</li>
      <li>Bấm <strong>"Tải lên mẫu âm thanh"</strong> và chọn file ghi âm của bạn.</li>
      <li>Đặt tên cho hồ sơ giọng (ví dụ: "Giọng của tôi") và bấm <strong>"Tạo hồ sơ giọng"</strong>.</li>
      <li>Khi tạo giọng đọc, chọn hồ sơ này là AI sẽ đọc bất kỳ văn bản nào bằng chính chất giọng của bạn!</li>
    </ol>

    <h2>5. Khắc Phục Sự Cố & Câu Hỏi Thường Gặp</h2>
    <div class="faq-item">
      <div class="faq-q">❓ Lỗi: "Missing packaged runtime-main Python..."</div>
      <p>👉 <strong>Cách xử lý:</strong> Do bạn chưa giải nén đủ 3 gói AI hoặc đặt sai vị trí thư mục. Đảm bảo thư mục <code>ai-packages</code> nằm cùng cấp với <code>Voca Basic.exe</code> và bên trong có đủ 3 thư mục con: <code>runtime-main</code>, <code>runtime-chatterbox</code>, <code>models</code>.</p>
    </div>

    <div class="faq-item">
      <div class="faq-q">❓ Làm sao để reset ứng dụng về bản trắng sạch sẽ?</div>
      <p>👉 <strong>Cách xử lý:</strong> Đóng ứng dụng và click đúp vào file <span class="code-tag">Reset-Data-Ban-Trang.bat</span> trong thư mục phần mềm. Toàn bộ lịch sử test và bản quyền sẽ được xóa sạch.</p>
    </div>

    <div class="faq-item">
      <div class="faq-q">❓ Phần mềm có cần mạng Internet để kiểm tra bản quyền không?</div>
      <p>👉 <strong>Trả lời:</strong> Hoàn toàn không. Bản quyền được xác thực mật mã Offline theo phần cứng máy tính.</p>
    </div>
  </div>

  <footer>
    <p>© 2026 Voca Basic. Chúc bạn có những trải nghiệm tuyệt vời cùng phần mềm!</p>
  </footer>
</div>

</body>
</html>
"""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(html, encoding="utf-8")
    print(f"Saved HTML: {out_path}")


def main():
    root = Path(__file__).resolve().parent.parent
    core_dist = root / "release" / "core-dist" / "Voca Basic Core"
    release_dir = root / "release"

    # Generate into core-dist (for customers)
    build_docx(core_dist / "Voca Basic - Huong Dan Su Dung.docx")
    build_txt(core_dist / "Huong-Dan-Su-Dung.txt")
    build_html(core_dist / "Huong-Dan-Su-Dung.html")

    # Also generate into release root (for seller convenience)
    build_docx(release_dir / "Voca Basic - Huong Dan Su Dung.docx")
    build_txt(release_dir / "Huong-Dan-Su-Dung.txt")
    build_html(release_dir / "Huong-Dan-Su-Dung.html")


if __name__ == "__main__":
    main()
